"""
Reference and rubric dataset loading.

Responsibilities:
- Load reference solutions from the canonical Excel repository.
- Load evaluation rubrics from JSON.
- Preserve backward-compatible helper functions used by
  evaluation/interviewer and evaluation/scoring.

Important architecture rule:
- Excel is the source of truth for reference solutions.
- This module does NOT select the target reference.
- This module does NOT use `Is Optimal Target`.
- Target selection/progression belongs to the adaptive-policy layer.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent

# The current repository keeps shared dataset assets under AI/.
AI_ROOT = PROJECT_ROOT / "AI"

DATASET_ROOT = AI_ROOT / "dataset"

REFERENCES_DIR = DATASET_ROOT / "references"

RUBRICS_DIR = DATASET_ROOT / "rubrics"

# Canonical reference-answer workbook.
REFERENCE_DATASET = (
    REFERENCES_DIR
    / "leetcode_1_to_10_reference_dataset_updated.xlsx"
)


# ============================================================
# JSON HELPERS
# ============================================================

def _load_json(path: Path) -> dict[str, Any]:
    """
    Load a JSON object from disk.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"Dataset file not found: {path}"
        )

    try:
        with path.open(
            "r",
            encoding="utf-8",
        ) as file:
            data = json.load(file)

    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Invalid JSON in dataset file: {path}"
        ) from exc

    if not isinstance(data, dict):
        raise TypeError(
            f"Dataset file must contain a JSON object: {path}"
        )

    return data


# ============================================================
# PROBLEM ID
# ============================================================

def _get_problem_id(problem: dict[str, Any]) -> str:
    """
    Extract the problem ID from a problem object.

    The problem object may come from:
    - GraphQL/API problem provider
    - existing tests
    - legacy callers

    The reference repository only needs Problem ID as the
    join key.
    """

    if not isinstance(problem, dict):
        raise TypeError(
            "problem must be a dictionary."
        )

    problem_id = (
        problem.get("problem_id")
        or problem.get("id")
    )

    if problem_id is None:
        raise ValueError(
            "Problem does not contain 'id' or 'problem_id'."
        )

    problem_id = str(problem_id).strip()

    if not problem_id:
        raise ValueError(
            "Problem ID cannot be empty."
        )

    return problem_id


# ============================================================
# REFERENCE DATASET
# ============================================================

def load_reference_dataset() -> dict[str, dict[str, Any]]:
    """
    Load the complete reference repository from Excel.

    Returns:

        {
            "P001": {
                "problem_id": "P001",
                "reference_solutions": [
                    {...},
                    {...},
                    ...
                ]
            },
            ...
        }

    Important:
    - Problem ID is used only as the join key.
    - Reference-specific information comes from Excel.
    - No target/optimal decision is made here.
    """

    if not REFERENCE_DATASET.exists():
        raise FileNotFoundError(
            "Reference dataset not found at: "
            f"{REFERENCE_DATASET}"
        )

    # Import locally so importing this module does not require
    # openpyxl unless the Excel repository is actually used.
    from openpyxl import load_workbook

    workbook = load_workbook(
        REFERENCE_DATASET,
        read_only=True,
        data_only=True,
    )

    try:
        if "Reference Solutions" not in workbook.sheetnames:
            raise ValueError(
                "Excel dataset does not contain the "
                "'Reference Solutions' sheet."
            )

        sheet = workbook["Reference Solutions"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(
            rows,
            None,
        )

        if not headers:
            raise ValueError(
                "Reference dataset is empty."
            )

        headers = [
            str(header).strip()
            if header is not None
            else ""
            for header in headers
        ]

        required_columns = {
            "Problem ID",
            "Reference ID",
            "Expected Approach",
            "Detailed Explanation",
            "Pseudocode",
            "Expected Data Structures",
            "Time Complexity",
            "Space Complexity",
            "Reasoning Steps",
            "Edge Cases",
            "Solution Type",
            "Next Better Reference ID",
            "Optimization Goal",
        }

        missing_columns = (
            required_columns
            - set(headers)
        )

        if missing_columns:
            raise ValueError(
                "Reference dataset is missing required "
                f"columns: {sorted(missing_columns)}"
            )

        dataset: dict[str, dict[str, Any]] = {}

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            reference: dict[str, Any] = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                reference[header] = value

            # ------------------------------------------------
            # Problem ID
            # ------------------------------------------------

            raw_problem_id = reference.get(
                "Problem ID"
            )

            if raw_problem_id is None:
                continue

            problem_id = str(
                raw_problem_id
            ).strip()

            if not problem_id:
                continue

            reference["Problem ID"] = problem_id

            # ------------------------------------------------
            # Reference ID
            # ------------------------------------------------

            raw_reference_id = reference.get(
                "Reference ID"
            )

            if raw_reference_id is None:
                raise ValueError(
                    f"Reference for {problem_id} is missing "
                    "'Reference ID'."
                )

            reference_id = str(
                raw_reference_id
            ).strip()

            if not reference_id:
                raise ValueError(
                    f"Reference for {problem_id} has an "
                    "empty 'Reference ID'."
                )

            reference["Reference ID"] = reference_id

            # ------------------------------------------------
            # Solution Type
            # ------------------------------------------------

            solution_type = reference.get(
                "Solution Type"
            )

            if solution_type is not None:
                reference["Solution Type"] = str(
                    solution_type
                ).strip()

            # ------------------------------------------------
            # Next Better Reference ID
            #
            # This is metadata for the adaptive layer.
            # We load it but DO NOT use it to select a target.
            # ------------------------------------------------

            next_reference = reference.get(
                "Next Better Reference ID"
            )

            if next_reference is not None:
                next_reference = str(
                    next_reference
                ).strip()

                reference[
                    "Next Better Reference ID"
                ] = (
                    next_reference
                    if next_reference
                    else None
                )

            # ------------------------------------------------
            # Create problem entry
            # ------------------------------------------------

            if problem_id not in dataset:
                dataset[problem_id] = {
                    "problem_id": problem_id,
                    "reference_solutions": [],
                }

            dataset[
                problem_id
            ][
                "reference_solutions"
            ].append(
                reference
            )

        if not dataset:
            raise ValueError(
                "Reference dataset contains no valid "
                "reference solutions."
            )

        return dataset

    finally:
        workbook.close()


# ============================================================
# REFERENCE SOLUTIONS
# ============================================================

def load_reference_solution(
    problem: dict[str, Any],
) -> list[dict[str, Any]]:
    """
    Load every reference solution associated with a problem.

    This function intentionally returns ALL references.

    It does NOT:
    - choose an optimal reference
    - choose a target
    - choose a next reference
    - inspect Is Optimal Target

    Those decisions belong to the adaptive-policy layer.
    """

    problem_id = _get_problem_id(
        problem
    )

    dataset = load_reference_dataset()

    problem_data = dataset.get(
        problem_id
    )

    if problem_data is None:
        raise FileNotFoundError(
            "No reference solutions found for "
            f"problem: {problem_id}"
        )

    references = problem_data.get(
        "reference_solutions",
        []
    )

    if not references:
        raise ValueError(
            f"Problem {problem_id} has no reference solutions."
        )

    return references


# ============================================================
# RUBRIC
# ============================================================

def load_rubric(
    problem: dict[str, Any],
) -> dict[str, Any]:
    """
    Load the evaluation rubric for a problem.

    Rubrics remain separate from the reference repository.
    """

    problem_id = _get_problem_id(
        problem
    )

    path = (
        RUBRICS_DIR
        / f"{problem_id}_rubric.json"
    )

    return _load_json(
        path
    )


# ============================================================
# EVALUATION CONTEXT
# ============================================================

def load_evaluation_context(
    problem: dict[str, Any],
) -> tuple[
    list[dict[str, Any]],
    dict[str, Any],
]:
    """
    Backward-compatible helper used by the existing evaluation
    and interviewer code.

    Returns:

        (
            reference_solutions,
            rubric
        )
    """

    reference_solutions = (
        load_reference_solution(
            problem
        )
    )

    rubric = load_rubric(
        problem
    )

    return (
        reference_solutions,
        rubric,
    )